"""Build an editable scientific DOCX from the reviewed chapter sources.

Uses python-docx for the style reference and deterministic OOXML finishing,
and Pandoc for Markdown/LaTeX to native Word math. No research is recalculated.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import subprocess
from pathlib import Path
from urllib.parse import unquote

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
REPO = ROOT.parents[3]
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
M = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
CACHED_FIGURES = {}
UNAVAILABLE_REFERENCES = []
ALLOW_UNAVAILABLE_REFERENCES = False
LINK_BASE = ROOT
CHAPTERS = [
    '01_literature.md', '02_alternatives.md', '03_patents.md',
    '04_problem.md', '05_anatomy.md', '06_forward.md',
    '07_experiments.md', '08_tissue.md', '09_heart_dynamics.md',
    '10_chambers.md', '11_validation.md', '12_discussion.md',
]


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def font(style, size: float, bold=False):
    style.font.name = 'Times New Roman'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        fonts.set(qn('w:' + key), 'Times New Roman')
    for key in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
        fonts.attrib.pop(qn('w:' + key), None)


def reference_docx(path: Path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin = Cm(3), Cm(1.5)
    sec.top_margin = sec.bottom_margin = Cm(2)
    sec.header_distance = sec.footer_distance = Cm(1.2)
    sec.different_first_page_header_footer = True
    for name in ['Normal', 'Body Text', 'First Paragraph']:
        if name not in doc.styles:
            doc.styles.add_style(name, 1)
        font(doc.styles[name], 14)
        pf = doc.styles[name].paragraph_format
        pf.line_spacing = 1.5
        pf.space_after = Pt(0)
        pf.first_line_indent = Cm(1.25)
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.widow_control = True
    for name in ['Heading 1', 'Heading 2', 'Heading 3']:
        font(doc.styles[name], 14, True)
        pf = doc.styles[name].paragraph_format
        pf.first_line_indent = Cm(0)
        pf.space_before, pf.space_after = Pt(16), Pt(8)
        pf.line_spacing = 1.15
        pf.keep_with_next = True
        pf.keep_together = True
        pf.page_break_before = name == 'Heading 1'
    font(doc.styles['Title'], 20, True)
    doc.styles['Title'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(doc.styles['Subtitle'], 14)
    doc.styles['Subtitle'].font.italic = False
    doc.styles['Subtitle'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(doc.styles['Caption'], 12)
    cp = doc.styles['Caption'].paragraph_format
    cp.first_line_indent = Cm(0)
    cp.line_spacing = 1.0
    cp.space_before, cp.space_after = Pt(5), Pt(9)
    cp.keep_together = True
    for name in ['TOC Heading', 'TOC 1', 'TOC 2', 'Bibliography']:
        if name not in doc.styles:
            doc.styles.add_style(name, 1)
        font(doc.styles[name], 12 if name != 'TOC Heading' else 14, name == 'TOC Heading')
        doc.styles[name].paragraph_format.line_spacing = 1.05
    for name in ['List Paragraph', 'List Bullet', 'List Number']:
        font(doc.styles[name], 14)
        doc.styles[name].paragraph_format.line_spacing = 1.25
        doc.styles[name].paragraph_format.space_after = Pt(3)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    doc.settings.element.append(OxmlElement('w:updateFields'))
    doc.settings.element[-1].set(qn('w:val'), 'true')
    lang = OxmlElement('w:lang')
    lang.set(qn('w:val'), 'ru-RU')
    doc.styles['Normal'].element.get_or_add_rPr().append(lang)
    doc.save(path)


def prepare_source(path: Path, assets: list, bibliography: list) -> str:
    text = path.read_text(encoding='utf-8-sig')
    text = re.sub(r'\\\((.*?)\\\)', lambda m: '$' + m.group(1) + '$', text, flags=re.S)
    text = re.sub(r'\\\[(.*?)\\\]', lambda m: '\n$$\n' + m.group(1) + '\n$$\n', text, flags=re.S)
    # The source remains collapsible; the printed report always exposes tables.
    text = re.sub(r'<summary>(.*?)</summary>', r'\n\n\1\n\n', text, flags=re.S)
    text = re.sub(r'</?details\b[^>]*>', '', text)
    match = re.search(r'^## Источники к главе[^\n]*\n', text, re.M)
    refs = ''
    if match:
        refs = text[match.end():].strip()
        text = text[:match.start()].rstrip()

    def image_link(match):
        alt, target = match.group(1), unquote(match.group(2).strip('<>'))
        source = (path.parent / target).resolve()
        origin = source.relative_to(REPO).as_posix() if source.is_relative_to(REPO) else source.name
        cache_used = False
        if not source.is_file():
            cached = CACHED_FIGURES.get(origin)
            if not cached:
                raise FileNotFoundError(f'Missing figure in {path.name}: {target}')
            source = ROOT / cached['file']
            if not source.is_file() or sha(source) != cached['sha256']:
                raise ValueError(f'Unverified cached figure: {target}')
            cache_used = True
        name = sha(source)[:12] + source.suffix.lower()
        dest = ROOT / 'figures' / name
        if not dest.exists():
            shutil.copy2(source, dest)
        assets.append({'chapter': path.name, 'file': 'figures/' + name,
                       'source': origin, 'sha256': sha(source), 'description': alt,
                       'verified_previous_build_cache': cache_used})
        return f'![{alt}](figures/{name})'

    text = re.sub(r'!\[([^\]]*)\]\(([^)]+)\)', image_link, text)

    def ordinary_link(match):
        label, target = match.group(1), match.group(2)
        if re.match(r'(https?://|mailto:|#)', target):
            return match.group(0)
        source = (path.parent / unquote(target.strip('<>').split('#')[0])).resolve()
        if not source.exists():
            if not ALLOW_UNAVAILABLE_REFERENCES:
                raise FileNotFoundError(f'Missing reference in {path.name}: {target}')
            UNAVAILABLE_REFERENCES.append({'chapter': path.name, 'label': label, 'target': target,
                                           'status': 'not_in_current_editorial_delivery'})
            return label + ' (материал не приложен)'
        try:
            relative = source.relative_to(REPO).as_posix()
        except ValueError:
            # Private absolute roots must not enter the reader document.
            return label
        relative = Path(os.path.relpath(source, LINK_BASE)).as_posix()
        return f'[{label}](<{relative}>)'

    text = re.sub(r'(?<!!)\[([^\]]+)\]\(([^)]+)\)', ordinary_link, text)
    refs = re.sub(r'(?<!!)\[([^\]]+)\]\(([^)]+)\)', ordinary_link, refs)
    bibliography.extend(x for x in re.split(r'\n(?=- |\d+\. )', refs)
                        if re.match(r'^(?:- |\d+\. )', x.strip()))
    def reflow_math(match):
        body = match.group(1).strip()
        # Explicit page-review decisions: break between complete equalities.
        reviewed_breaks = {'7.4', '8.1', '8.5', '8.14', '8.21', '9.1',
                           '9.8', '9.9', '9.11', '9.14', '10.1', '10.2',
                           '10.3', '10.19', '11.5'}
        force_split = bool(set(re.findall(r'\\tag\{([^}]+)\}', body)) & reviewed_breaks)
        needs_split = force_split or len(body) >= 240 or (len(body) >= 180 and body.count(r'\qquad') >= 2)
        separator = r'\\q{1,2}uad' if force_split else r'\\qquad'
        if not needs_split or not re.search(separator, body) or r'\begin' in body:
            return match.group(0)
        tags = re.findall(r'\\tag\{[^}]+\}', body)
        body = re.sub(r'\\tag\{[^}]+\}', '', body).strip()
        # Split only complete clauses, outside braces and paired delimiters.
        depth = left = 0
        cuts = []
        for token in re.finditer(r'\\left|\\right|' + separator + r'|(?<!\\)[{}]', body):
            t = token.group()
            if t == '{': depth += 1
            elif t == '}': depth -= 1
            elif t == r'\left': left += 1
            elif t == r'\right': left -= 1
            elif depth == 0 and left == 0: cuts.append((token.start(), token.end()))
        if not cuts:
            return match.group(0)
        pieces, start = [], 0
        for a, b in cuts:
            pieces.append(body[start:a].strip()); start = b
        pieces.append(body[start:].strip())
        return '\n$$\n\\begin{gathered}\n' + '\\\\\n'.join(pieces) + '\n\\end{gathered}\n' + ''.join(tags) + '\n$$\n'
    text = re.sub(r'\$\$(.*?)\$\$', reflow_math, text, flags=re.S)
    # Pandoc's texmath does not consistently support \tag in OMML.
    # Keep the number inside native math as \qquad\text{(N.k)}.
    text = re.sub(r'\\tag\{([^}]+)\}', r'\\qquad\\text{(\1)}', text)
    return text.strip()


def finish_docx(path: Path):
    doc = Document(path)
    in_bibliography = False
    for p in doc.paragraphs:
        text = p.text.strip()
        if p.style.name == 'Heading 1':
            in_bibliography = text == 'Список использованных источников'
        elif in_bibliography:
            p.style = doc.styles['Bibliography']
            p.paragraph_format.left_indent = Cm(0.95)
            p.paragraph_format.first_line_indent = Cm(-0.95)
            p.paragraph_format.tab_stops.clear_all()
            p.paragraph_format.tab_stops.add_tab_stop(Cm(0.95))
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = 1.1
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.size = Pt(12)
        if text == '[[TOC]]':
            p.clear()
            p.paragraph_format.first_line_indent = Cm(0)
            field = OxmlElement('w:fldSimple')
            field.set(qn('w:instr'), 'TOC \\o "1-2" \\h \\z \\u')
            p._p.append(field)
        elif text in ['Содержание', 'Реферат']:
            p.style = doc.styles['TOC Heading']
            p.paragraph_format.page_break_before = True
        elif re.match(r'^Рисунок\s+\d+\.\d+\s+[—–-]', text):
            p.style = doc.styles['Caption']
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif re.match(r'^Таблица\s+[\dА-ЯA-Z]+\.\d+\s+[—–-]', text):
            p.style = doc.styles['Caption']
            p.paragraph_format.keep_with_next = True
        if p._p.xpath('.//w:drawing'):
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
        if p._p.xpath('.//m:oMathPara'):
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.1
        following = p._p.getnext()
        if following is not None and next(following.iter(qn('m:oMathPara')), None) is not None and (text.endswith(':') or re.search(r'(?:вид|как|выражением|формулами|записи|равенством)\s*\.?$', text)):
            p.paragraph_format.keep_with_next = True
    for shape in doc.inline_shapes:
        scale = min(Cm(16.3) / shape.width, Cm(18) / shape.height, 1)
        shape.width, shape.height = int(shape.width * scale), int(shape.height * scale)
    for table in doc.tables:
        table.autofit = False
        table.style = 'Table Grid'
        properties = table._tbl.tblPr
        borders = OxmlElement('w:tblBorders')
        for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            border = OxmlElement('w:' + edge)
            for key, value in [('val', 'single'), ('sz', '4'), ('color', 'B7B7B7')]:
                border.set(qn('w:' + key), value)
            borders.append(border)
        properties.append(borders)
        margins = OxmlElement('w:tblCellMar')
        for edge in ['top', 'left', 'bottom', 'right']:
            margin = OxmlElement('w:' + edge)
            margin.set(qn('w:w'), '70')
            margin.set(qn('w:type'), 'dxa')
            margins.append(margin)
        properties.append(margins)
        n = len(table.columns)
        # Length weighting improves text-heavy columns without crushing numbers.
        weights = []
        for c in range(n):
            lengths = [len(row.cells[c].text) for row in table.rows]
            weights.append(max(8, min(80, sum(lengths) / len(lengths))) ** 0.5)
        widths = [16.5 * x / sum(weights) for x in weights]
        if n == 4 and table.cell(0, 0).text.strip() == 'Том':
            # Keep archive identifiers readable without splitting a word mid-letter.
            widths = [3.7, 2.8, 2.8, 7.2]
        if n == 4 and table.cell(0, 0).text.strip() == 'Код':
            widths = [1.15, 5.0, 4.1, 6.25]
        for col, width in zip(table.columns, widths):
            col.width = Cm(width)
        for ri, row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            if ri == 0:
                row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for ci, cell in enumerate(row.cells):
                cell.width = Cm(widths[ci])
                for math_run in cell._tc.xpath('.//m:r'):
                    properties = math_run.find(qn('w:rPr'))
                    if properties is None:
                        properties = OxmlElement('w:rPr')
                        math_run.insert(0, properties)
                    size = properties.find(qn('w:sz'))
                    if size is None:
                        size = OxmlElement('w:sz'); properties.append(size)
                    size.set(qn('w:val'), '22')
                for p in cell.paragraphs:
                    pf = p.paragraph_format
                    pf.first_line_indent = Cm(0)
                    pf.line_spacing = 1.0
                    pf.space_before = pf.space_after = Pt(3)
                    pf.keep_with_next = ri == 0
                    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for run in p.runs:
                        run.font.name, run.font.size = 'Times New Roman', Pt(11)
                        if ri == 0:
                            run.font.bold = True
                if ri == 0:
                    shd = OxmlElement('w:shd')
                    shd.set(qn('w:fill'), 'EDEDED')
                    cell._tc.get_or_add_tcPr().append(shd)
    # A simple title page; unknown legal and personal details are omitted.
    title_lines = [
        ('Отчёт ГПХ РНЦХ', 'Title'),
        ('Разработка методики определения ударного объёма по электроимпедансным измерениям', 'Title'),
        ('Индивидуальная анатомия и динамика сердца', 'Subtitle'),
        ('Научная редакция по состоянию на 29 октября 2026 года', 'Subtitle'),
        ('Редакция с матрицей соответствия договору', 'Subtitle'),
        ('Этап разработки и проверки методики', 'Subtitle'),
    ]
    first = doc.element.body[0]
    for text, style in title_lines:
        p = doc.add_paragraph(text, style)
        p.paragraph_format.space_before = Pt(24)
        p.paragraph_format.space_after = Pt(18)
        first.addprevious(p._p)
    doc.core_properties.title = title_lines[1][0]
    doc.core_properties.subject = 'Отчёт о разработке методики'
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.core_properties.comments = ''
    # Word/Pandoc templates may carry a blue rule on the Title style.
    for root in [doc.styles.element, doc.element]:
        for node in root.xpath('.//w:pBdr'):
            node.getparent().remove(node)
    doc.save(path)


def main():
    global ALLOW_UNAVAILABLE_REFERENCES, LINK_BASE
    ap = argparse.ArgumentParser()
    ap.add_argument('--output', type=Path, required=True)
    ap.add_argument('--pandoc', default='pandoc')
    ap.add_argument('--partial', action='store_true')
    ap.add_argument('--annotate-unavailable-references', action='store_true',
                    help='Explicitly mark historical references absent from this delivery; record them in the manifest.')
    ap.add_argument('--figure-cache-manifest', type=Path,
                    help='Previous build receipt, used only after matching each figure SHA-256.')
    args = ap.parse_args()
    LINK_BASE = args.output.resolve().parent
    ALLOW_UNAVAILABLE_REFERENCES = args.annotate_unavailable_references
    if args.figure_cache_manifest:
        cache = json.loads(args.figure_cache_manifest.read_text(encoding='utf-8-sig'))
        CACHED_FIGURES.update({row['source']: row for row in cache['figures']})
    args.output.parent.mkdir(parents=True, exist_ok=True)
    assets, bibliography, parts, source_manifest, citation_sets = [], [], [], [], []
    paths = [ROOT / 'front_matter.md'] + [ROOT / 'chapters' / n for n in CHAPTERS] + [ROOT / 'conclusion.md']
    for path in paths:
        if not path.exists() and args.partial:
            continue
        if not path.exists():
            raise FileNotFoundError(path)
        first_ref = len(bibliography)
        parts.append(prepare_source(path, assets, bibliography))
        citation_sets.append(bibliography[first_ref:])
        source_manifest.append({'file': path.relative_to(ROOT).as_posix(), 'sha256': sha(path)})
    def reference_key(ref):
        dois = re.findall(r'https?://(?:dx\.)?doi\.org/([^\s)<>]+)', ref)
        urls = re.findall(r'\]\(([^)]+)\)', ref)
        if len(set(urls)) > 1:
            # A compound project reference may cite several different files.
            # Deduplicating by its first URL would silently discard the rest.
            return 'compound:' + '|'.join(sorted(set(u.casefold() for u in urls)))
        return (dois[0].lower() if dois else urls[0].casefold() if urls
                else re.sub(r'\W+', '', re.sub(r'^\d+\.\s*', '', ref)).lower())
    seen, references, reference_numbers = set(), [], {}
    for ref in bibliography:
        if not ref.strip():
            continue
        key = reference_key(ref)
        if key not in seen:
            seen.add(key)
            clean = re.sub(r'^(?:\d+\.|-)\s+', '', ref.strip())
            references.append(f'{len(references) + 1}. {clean}')
            reference_numbers[key] = len(references)
    if references:
        citation_map = []
        # Chapter-local [1–7] citations are renumbered before URL citations.
        for index, local_refs in enumerate(citation_sets):
            local_numbers = {}
            for ref in local_refs:
                m = re.match(r'^(\d+)\.', ref.strip())
                if m:
                    local_numbers[int(m.group(1))] = reference_numbers[reference_key(ref)]
            citation_map.append({'source': source_manifest[index]['file'],
                                 'local_to_global': local_numbers})
            def renumber(m):
                ids = []
                for item in re.split(r'[,;]\s*', m.group(1)):
                    ends = re.split(r'[–—-]', item.strip())
                    ids.extend(range(int(ends[0]), int(ends[-1])+1))
                if not ids or not all(n in local_numbers for n in ids):
                    return m.group(0)
                mapped = [local_numbers[n] for n in ids]
                if len(mapped) > 2 and mapped == list(range(mapped[0], mapped[-1]+1)):
                    return f'[{mapped[0]}–{mapped[-1]}]'
                return '[' + ', '.join(map(str, mapped)) + ']'
            chunks = re.split(r'(\$\$.*?\$\$|\$[^$\n]+\$)', parts[index], flags=re.S)
            for i in range(0, len(chunks), 2):
                chunks[i] = re.sub(r'(?<![!\\])\[([0-9]+(?:\s*[,;–—-]\s*[0-9]+)*)\](?!\()', renumber, chunks[i])
            parts[index] = ''.join(chunks)
        # Preserve a readable citation even when a URL alias has no exact match.
        url_index = {}
        doi_index = {}
        for number, ref in enumerate(references, 1):
            for url in re.findall(r'https?://[^\s)<>]+', ref):
                url_index[url.rstrip('/')] = number
            for doi in re.findall(r'10\.\d{4,9}/[^\s)<>]+', ref):
                doi_index[doi.lower().rstrip('.')] = number
        def citation(m):
            label, url = m.group(1), m.group(2)
            number = url_index.get(url.rstrip('/'))
            if number is None:
                doi = re.search(r'10\.\d{4,9}/[^\s)<>]+', url)
                if doi:
                    number = doi_index.get(doi.group().lower().rstrip('.'))
            pinpoint = re.search(r'(?:§|с\.|табл\.|рис\.|пункты?\b).*', label)
            visible = (str(number) + (', ' + pinpoint.group() if pinpoint else '')) if number else label
            return f'[\\[{visible}\\]]({url})'
        parts = [re.sub(r'(?<!!)\[([^\]]+)\]\((https?://[^)]+)\)', citation, part)
                 for part in parts]
        parts.append('# Список использованных источников\n\n' + '\n\n'.join(references))
        (ROOT / 'bibliography_map.json').write_text(json.dumps(
            {'chapters': citation_map, 'references': references}, ensure_ascii=False, indent=2), encoding='utf-8')
    appendix = ROOT / 'appendices.md'
    if appendix.exists():
        parts.append(prepare_source(appendix, assets, bibliography))
        source_manifest.append({'file': appendix.relative_to(ROOT).as_posix(), 'sha256': sha(appendix)})
    combined = '\n\n'.join(parts) + '\n'
    source = ROOT / 'report_for_word.md'
    source.write_text(combined, encoding='utf-8')
    ref = args.output.parent / 'reference_styles.docx'
    reference_docx(ref)
    command = [args.pandoc, str(source), '-f', 'markdown+tex_math_dollars-implicit_figures',
               '--resource-path', str(ROOT), '--reference-doc', str(ref),
               '-o', str(args.output)]
    result = subprocess.run(command, capture_output=True, text=True, encoding='utf-8', errors='replace')
    (args.output.parent / 'pandoc.log').write_text(result.stderr, encoding='utf-8')
    if result.returncode:
        raise RuntimeError(result.stderr)
    if 'Could not convert TeX math' in result.stderr or 'Could not fetch resource' in result.stderr:
        raise RuntimeError('Pandoc reported unsupported math or missing resources; inspect pandoc.log')
    finish_docx(args.output)
    receipt = {'sources': source_manifest, 'figures': assets, 'bibliography_entries': len(references),
               'document_sha256_before_fields': sha(args.output), 'partial_build': args.partial,
               'unavailable_references': UNAVAILABLE_REFERENCES,
               'unavailable_reference_policy': 'annotate' if ALLOW_UNAVAILABLE_REFERENCES else 'strict',
               'local_links_relative_to_docx_directory': True,
               'figure_cache_manifest': args.figure_cache_manifest.name if args.figure_cache_manifest else None}
    (ROOT / 'build_manifest.json').write_text(json.dumps(receipt, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'output': str(args.output), 'sources': len(source_manifest),
                      'figures': len(assets), 'references': len(references)}, ensure_ascii=False))


if __name__ == '__main__':
    main()
