from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parent.parent
INPUT = ROOT / "07_Литературный_обзор.docx"
OUTPUT = ROOT / ".evidence_work" / "07_Литературный_обзор.formatted.docx"


def set_font(run, size: float, bold: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def configure_style(style, size: float, *, bold: bool = False, first_line: bool = True) -> None:
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    fmt = style.paragraph_format
    fmt.line_spacing = 1.5
    fmt.space_after = Pt(0)
    fmt.space_before = Pt(0)
    fmt.first_line_indent = Cm(1.25) if first_line else Cm(0)
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def add_page_number(section) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])
    set_font(run, 12)


def main() -> None:
    doc = Document(INPUT)

    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(3)
        section.right_margin = Cm(1.5)
        section.header_distance = Cm(1.25)
        section.footer_distance = Cm(1.25)
        add_page_number(section)

    for name in ("Normal", "Body Text"):
        if name in doc.styles:
            configure_style(doc.styles[name], 14, first_line=True)
    if "List Paragraph" in doc.styles:
        configure_style(doc.styles["List Paragraph"], 14, first_line=False)
        doc.styles["List Paragraph"].paragraph_format.left_indent = Cm(1.25)
    for index in range(1, 4):
        name = f"Heading {index}"
        if name not in doc.styles:
            continue
        configure_style(doc.styles[name], 14, bold=True, first_line=False)
        style = doc.styles[name]
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if index == 1 else 8)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER if index == 1 else WD_ALIGN_PARAGRAPH.LEFT
        )
    for name in ("TOC Heading", "Contents Heading"):
        if name in doc.styles:
            configure_style(doc.styles[name], 14, bold=True, first_line=False)
            doc.styles[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for index in range(1, 10):
        name = f"TOC {index}"
        if name in doc.styles:
            configure_style(doc.styles[name], 12, first_line=False)
            doc.styles[name].paragraph_format.line_spacing = 1.0

    for paragraph in doc.paragraphs:
        style_name = paragraph.style.name if paragraph.style else ""
        if style_name.startswith("Heading") or style_name.startswith("TOC"):
            paragraph.paragraph_format.first_line_indent = Cm(0)
        for run in paragraph.runs:
            size = 12 if style_name.startswith("TOC") else 14
            set_font(run, size)

    for table in doc.tables:
        table.autofit = True
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.first_line_indent = Cm(0)
                    paragraph.paragraph_format.line_spacing = 1.0
                    paragraph.paragraph_format.space_after = Pt(2)
                    for run in paragraph.runs:
                        set_font(run, 10, bold=True if row_index == 0 else None)

    core = doc.core_properties
    core.title = "Современные методы мониторинга центральной гемодинамики и возможности прекардиальной импедансной кардиографии"
    core.subject = "Тематический литературный обзор фиксированного корпуса Zotero"
    core.keywords = "гемодинамика; импедансная кардиография; сердечный выброс; прекардиальное картирование"

    doc.save(OUTPUT)
    OUTPUT.replace(INPUT)
    print(INPUT)


if __name__ == "__main__":
    main()
