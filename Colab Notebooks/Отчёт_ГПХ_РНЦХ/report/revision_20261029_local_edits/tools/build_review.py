"""Assemble the new scientific report with native Word math and cross-references."""
from __future__ import annotations
import argparse, copy, hashlib, importlib.util, json, re, subprocess
from pathlib import Path
from collections import Counter
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from lxml import etree

ROOT=Path(__file__).resolve().parents[1]
REPO=ROOT.parents[3]
OLD=ROOT.parent/'revision_20261028'
PANDOC=Path('J:/PC/Pandoc/pandoc.exe')
ORDER=['00_front.md','01_literature.md','02_materials.md','03_models.md','04_processing.md','05_model_results.md','06_tissue_results.md','07_stroke_volume.md','08_discussion.md','09_conclusion.md']
OUT=ROOT/'deliverables/Отчёт_ГПХ_РНЦХ_20261029_точечная_правка.docx'
FIGURES={}; BOOKMARK_ID=100

def jread(p): return json.loads(p.read_text(encoding='utf-8-sig'))
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def bmname(kind,key): return kind+'_'+hashlib.sha1(key.encode()).hexdigest()[:18]

def field(parent,instruction,value,bookmark=None):
    global BOOKMARK_ID
    if bookmark:
        BOOKMARK_ID+=1; start=OxmlElement('w:bookmarkStart'); start.set(qn('w:id'),str(BOOKMARK_ID)); start.set(qn('w:name'),bookmark); parent.append(start)
    el=OxmlElement('w:fldSimple'); el.set(qn('w:instr'),instruction)
    run=OxmlElement('w:r'); txt=OxmlElement('w:t'); txt.text=str(value); run.append(txt); el.append(run); parent.append(el)
    if bookmark:
        end=OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'),str(BOOKMARK_ID)); parent.append(end)

def text_run(parent,s):
    if not s: return
    r=OxmlElement('w:r'); t=OxmlElement('w:t'); t.set(qn('xml:space'),'preserve'); t.text=s; r.append(t); parent.append(r)

def load_figures():
    path=ROOT/'figures/manifest.json'
    if not path.exists(): return {}
    data=jread(path)
    if isinstance(data,dict): data=data.get('figures',data.get('items',data))
    if isinstance(data,dict): data=[dict(v,id=k) for k,v in data.items()]
    result={}
    for f in data:
        result[f['id']]=f
        for a in f.get('aliases',[]):
            selected=copy.deepcopy(f)
            ap=f.get('alias_paths',{}).get(a)
            if ap:
                selected['path']=ap
                matching=[x for x in f.get('assets',[]) if x.get('path')==ap]
                if matching: selected.update(copy.deepcopy(matching[-1]))
            result[a]=selected
    aliases_path=ROOT/'figures/aliases.json'
    if aliases_path.exists():
        for k,v in jread(aliases_path).items(): result[k]=result[v]
    overrides=ROOT/'figures/report_overrides.json'
    if overrides.exists():
        for f in jread(overrides):
            result[f['id']]=f
            for a in f.get('aliases',[]): result[a]=f
    return result

def assemble():
    global FIGURES
    FIGURES=load_figures(); bibliography=jread(ROOT/'bibliography.json')
    paths=[ROOT/'drafts'/name for name in ORDER]
    missing=[str(p) for p in paths if not p.exists()]
    if missing: raise RuntimeError('Missing chapters: '+', '.join(missing))
    parts=[p.read_text(encoding='utf-8-sig') for p in paths]
    # Drafts supply a caption proposal after each marker. Use the image curator's
    # source-checked caption once, rather than printing two captions per figure.
    def remove_caption_proposal(t):
        return re.sub(r'(\[\[FIG:[^\]]+\]\])\s*\n\s*\n([^\n]+)(?=\n|$)',r'\1\n',t)
    parts=[remove_caption_proposal(t) for t in parts]
    core='\n\n'.join(parts)
    apppaths=[p for p in [ROOT/'drafts/appendix_models.md',ROOT/'drafts/appendix_cardiac.md'] if p.exists()]
    apps='\n\n'.join(p.read_text(encoding='utf-8-sig') for p in apppaths)
    # Collect bibliography in first-use order across main text and scientific appendices.
    keys=[]
    for m in re.finditer(r'\[\[C:([^\]|]+)(?:\|([^\]]+))?\]\]',core+'\n'+apps):
        k=m[1]; k=bibliography.get(k,{}).get('alias',k)
        if k not in bibliography: raise RuntimeError('Unknown scientific reference '+k)
        if k not in keys: keys.append(k)
    nums={k:i+1 for i,k in enumerate(keys)}
    def cite(m):
        k=m[1]; k=bibliography[k].get('alias',k)
        return '['+str(nums[k])+(', '+m[2] if m[2] else '')+']'
    core=re.sub(r'\[\[C:([^\]|]+)(?:\|([^\]]+))?\]\]',cite,core)
    apps=re.sub(r'\[\[C:([^\]|]+)(?:\|([^\]]+))?\]\]',cite,apps)
    bib='# Список научной литературы\n\n'+'\n\n'.join(f'{nums[k]}. {bibliography[k]["text"]}' for k in keys)
    patents=ROOT/'drafts/patents.md'
    complete=core+'\n\n'+bib+'\n\n'+(patents.read_text(encoding='utf-8') if patents.exists() else '')+'\n\n'+apps
    counts=Counter(); labels={}; figure_manifest=[]
    def fig(m):
        key=m[1]
        if key not in FIGURES: raise RuntimeError('Missing figure '+key)
        f=FIGURES[key]; counts['fig']+=1; n=counts['fig']; labels[('fig',key)]=n
        src=Path(f.get('path',f.get('file','')))
        if not src.is_absolute(): src=(ROOT/src) if (ROOT/src).exists() else (REPO/src)
        if not src.exists(): raise RuntimeError('Missing figure file '+str(src))
        caption=f.get('caption',key)
        caption=re.sub(r'^Рисунок\s+[\w.]+\s*[—–-]\s*','',caption)
        # Explicit width keeps all diagrams within the printable area.
        from PIL import Image
        w,h=Image.open(src).size
        maxheight=18.0 if f['id']=='ct_segmentation' else 11.5
        crop=f.get('word_crop',{})
        ratio=(w*(1-crop.get('left',0)-crop.get('right',0)))/(h*(1-crop.get('top',0)-crop.get('bottom',0)))
        width=min(16.5,maxheight*ratio)
        figure_manifest.append(dict(key=key,number=n,path=str(src),sha256=sha(src),source=f.get('source'),caption=caption,word_crop=crop,width_cm=width,visible_height_cm=width/ratio))
        return f'\n\n![](<{src.as_posix()}>){{width={width:.3f}cm}}\n\nРисунок XSEQfig{n}X — {caption}\n\n'
    complete=re.sub(r'\[\[FIG:([^\]]+)\]\]',fig,complete)
    def table(m):
        key=m[1]; counts['tab']+=1; n=counts['tab']; labels[('tab',key)]=n
        return f'Таблица XSEQtab{n}X — {m[2]}'
    complete=re.sub(r'\[\[TABLE:([^\]|]+)\|([^\]]+)\]\]',table,complete)
    def equation(m):
        key=m[1]; counts['eq']+=1; n=counts['eq']; labels[('eq',key)]=n
        return f'XEQNUM{n}X'
    complete=re.sub(r'\[\[EQ:([^\]]+)\]\]',equation,complete)
    ref_usage=[]
    def ref(m):
        kind,key=m[1],m[2]
        if (kind,key) not in labels: raise RuntimeError('Unresolved reference '+kind+':'+key)
        n=labels[(kind,key)]; ref_usage.append((kind,key,n)); return f'XREF{kind}{n}X'
    complete=re.sub(r'\{\{(fig|tab|eq):([^}]+)\}\}',ref,complete)
    unresolved=re.findall(r'\[\[(?!TOC)[^\]]+\]\]|\{\{[^}]+\}\}',complete)
    if unresolved: raise RuntimeError('Unresolved markers '+repr(unresolved[:20]))
    complete=complete.replace('[[TOC]]','XTOCX')
    (ROOT/'report_for_word.md').write_text(complete,encoding='utf-8')
    manifest={'source_snapshot_sha256':jread(ROOT/'audit/source_manifest.json')['sha256'] if 'sha256' in jread(ROOT/'audit/source_manifest.json') else sha(ROOT/'audit/author_snapshot.docx'),
        'chapters':[{'path':str(p.relative_to(ROOT)),'sha256':sha(p)} for p in paths+apppaths],
        'main_words':len(re.findall(r'\S+',core)), 'figures':figure_manifest,'counts':dict(counts),'reference_keys':keys,'reference_numbers':nums,'cross_references':ref_usage,
        'labels':[{'kind':k[0],'key':k[1],'number':v} for k,v in labels.items()]}
    (ROOT/'build_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    return manifest

def finish(path,manifest):
    global BOOKMARK_ID
    doc=Document(path)
    # Pandoc already assigns bookmark IDs to headings. Allocate a disjoint
    # range; otherwise Word may repair starts and ends inconsistently.
    BOOKMARK_ID=max([0]+[int(e.get(qn('w:id'))) for e in doc.element.xpath('//w:bookmarkStart')])+1000
    toc_style=doc.styles.add_style('ReportContentsHeading',WD_STYLE_TYPE.PARAGRAPH)
    toc_style.base_style=doc.styles['Heading 1']
    outline=OxmlElement('w:outlineLvl');outline.set(qn('w:val'),'9')
    toc_style.element.get_or_add_pPr().append(outline)
    for shape,f in zip(doc.inline_shapes,manifest['figures']):
        crop=f.get('word_crop') or {}
        if crop:
            fill=shape._inline.xpath('.//pic:blipFill')[0]
            rect=OxmlElement('a:srcRect')
            for a,k in [('l','left'),('r','right'),('t','top'),('b','bottom')]: rect.set(a,str(round(100000*crop.get(k,0))))
            fill.insert(1,rect)
            shape.width=Cm(f['width_cm']); shape.height=Cm(f['visible_height_cm'])
    # A restrained title page matching the author's scientific report.
    first=doc.paragraphs[0]
    p=first.insert_paragraph_before('СПРАВКА-ОТЧЁТ',style='Title')
    p.paragraph_format.space_before=Pt(90); p.paragraph_format.space_after=Pt(40)
    p=first.insert_paragraph_before('Разработка методики определения ударного объёма сердца по электроимпедансным измерениям',style='Title')
    for r in p.runs: r.font.size=Pt(18)
    p=first.insert_paragraph_before('Научно-исследовательская работа',style='Subtitle'); p.paragraph_format.space_before=Pt(30)
    p=first.insert_paragraph_before('2026',style='Subtitle'); p.paragraph_format.space_before=Pt(110)
    labels={(x['kind'],x['number']):bmname(x['kind'],x['key']) for x in manifest['labels']}
    numbered={}
    paras=list(doc.paragraphs)
    for p in paras:
        text=p.text.strip()
        if text.startswith('ЭКГ обрабатывали последовательно'):
            p.paragraph_format.keep_together=True
        if text=='Содержание':p.style=toc_style
        if text=='XTOCX':
            p.clear(); p.paragraph_format.first_line_indent=Cm(0)
            field(p._p,' TOC \\o "1-2" \\h \\z \\u ','Содержание обновляется при открытии документа.')
        m=re.fullmatch(r'XEQNUM(\d+)X',text)
        if m:
            n=int(m[1]); prev=p._p.getprevious()
            if prev is None or not prev.xpath('.//m:oMathPara'): raise RuntimeError('Equation label without preceding display equation '+str(n))
            # Keep equation and its number in one native equation paragraph.
            text_run(prev,'  ('); field(prev,' SEQ Equation \\* ARABIC ',n,labels[('eq',n)]); text_run(prev,')')
            p._p.getparent().remove(p._p); numbered[('eq',n)]=True
            continue
        if p._p.xpath('.//w:drawing'):
            p.paragraph_format.first_line_indent=Cm(0); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=True; p.paragraph_format.space_before=Pt(7); p.paragraph_format.space_after=Pt(0)
        if text.startswith('Рисунок XSEQ') or text.startswith('Таблица XSEQ'):
            p.style=doc.styles['Caption']; p.paragraph_format.keep_with_next=text.startswith('Таблица')
            p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY if text.startswith('Рисунок') else WD_ALIGN_PARAGRAPH.LEFT
        if p._p.xpath('.//m:oMathPara'):
            p.paragraph_format.first_line_indent=Cm(0); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing=1.0; p.paragraph_format.space_before=Pt(6); p.paragraph_format.space_after=Pt(6)
            p.paragraph_format.keep_together=True
            intro=p._p.getprevious()
            if intro is not None and intro.tag==qn('w:p'):
                intro.get_or_add_pPr().append(OxmlElement('w:keepNext'))
    allparas=list(doc.paragraphs)+[p for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs]
    pat=re.compile(r'X(SEQ|REF)(fig|tab|eq)(\d+)X')
    for p in allparas:
        for run in list(p.runs):
            if not pat.search(run.text): continue
            old=run._r; container=OxmlElement('w:p'); pos=0
            for m in pat.finditer(run.text):
                text_run(container,run.text[pos:m.start()]); mode,kind,n=m[1],m[2],int(m[3])
                bookmark=labels[(kind,n)]
                if mode=='SEQ':
                    field(container,' SEQ '+{'fig':'Figure','tab':'Table','eq':'Equation'}[kind]+' \\* ARABIC ',n,bookmark); numbered[(kind,n)]=True
                else: field(container,' REF '+bookmark+' \\h ',n)
                pos=m.end()
            text_run(container,run.text[pos:])
            for child in list(container): old.addprevious(child)
            old.getparent().remove(old)
    # Scientific bibliography is numbered once, uses hanging indents and black text.
    in_bib=False; in_patents=False
    for p in doc.paragraphs:
        if p.text=='Список научной литературы': in_bib=True; continue
        if p.text=='Патентные материалы': in_patents=True; in_bib=False; continue
        if in_bib and p.style.name.startswith('Heading'): in_bib=False
        if in_patents and p.style.name.startswith('Heading'): in_patents=False
        if in_patents:p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        if in_bib:
            p.style=doc.styles['Bibliography']; p.paragraph_format.first_line_indent=Cm(-0.8); p.paragraph_format.left_indent=Cm(0.8)
            p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.0
    # Intentional widths based on content; repeated header, natural row heights.
    for table in doc.tables:
        table.autofit=False; table.style='Table Grid'; pr=table._tbl.tblPr
        borders=OxmlElement('w:tblBorders')
        for side in ['top','left','bottom','right','insideH','insideV']:
            b=OxmlElement('w:'+side); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'4'); b.set(qn('w:color'),'D9D9D9'); borders.append(b)
        prior=pr.find(qn('w:tblBorders'))
        if prior is not None: pr.remove(prior)
        pr.append(borders)
        margins=OxmlElement('w:tblCellMar')
        for side in ['top','bottom','left','right']:
            el=OxmlElement('w:'+side); el.set(qn('w:w'),'75'); el.set(qn('w:type'),'dxa'); margins.append(el)
        pr.append(margins)
        ncols=len(table.columns); weights=[]
        for c in range(ncols):
            lens=sorted(len(row.cells[c].text) for row in table.rows); q=lens[min(len(lens)-1,int(.7*len(lens)))]; weights.append(max(1.2,min(4.8,q**.48)))
        widths=[16.5*w/sum(weights) for w in weights]
        if ncols==5 and table.rows[0].cells[1].text=='Телосложение':widths=[1.2,3.8,3.8,4.0,3.7]
        if ncols==5 and table.rows[0].cells[0].text=='Представление':widths=[5.0,2.5,3.4,3.5,2.1]
        if ncols==4 and table.rows[0].cells[1].text=='Обратная форма':widths=[2.4,3.4,4.7,6.0]
        if ncols==4 and table.rows[0].cells[0].text=='Ткань или состояние':widths=[5.0,3.0,3.3,5.2]
        for col,w in zip(table.columns,widths): col.width=Cm(w)
        for ri,row in enumerate(table.rows):
            trpr=row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
            if ri==0: trpr.append(OxmlElement('w:tblHeader'))
            for ci,cell in enumerate(row.cells):
                cell.width=Cm(widths[ci]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if ri==0:
                    shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E8E8E8'); cell._tc.get_or_add_tcPr().append(shade)
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next=(ri==0)
                    p.paragraph_format.first_line_indent=Cm(0); p.paragraph_format.space_after=Pt(3); p.paragraph_format.space_before=Pt(3); p.paragraph_format.line_spacing=1.0
                    p.alignment=WD_ALIGN_PARAGRAPH.CENTER if len(p.text)<35 else WD_ALIGN_PARAGRAPH.LEFT
                    for r in p.runs: r.font.name='Times New Roman'; r.font.size=Pt(11); r.bold=(ri==0)
    # Ensure no title/theme decorations and no unexpected authorship metadata.
    for name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        st=doc.styles[name]; st.font.color.rgb=RGBColor(0,0,0)
        pr=st.element.find(qn('w:pPr'))
        if pr is not None:
            for el in list(pr):
                if el.tag in [qn('w:pBdr'),qn('w:shd')]: pr.remove(el)
    for name in ['Footnote Text','Footnote Reference']:
        if name in doc.styles:
            doc.styles[name].font.name='Times New Roman'
            doc.styles[name].font.size=Pt(10)
    if 'Footnote Text' in doc.styles:
        pf=doc.styles['Footnote Text'].paragraph_format
        pf.line_spacing=1.0; pf.first_line_indent=Cm(0)
        pf.space_before=Pt(0); pf.space_after=Pt(0)
    # Pandoc's footnote parts may carry direct formatting from the reference.
    for part in doc.part.related_parts.values():
        if part.content_type.endswith('wordprocessingml.footnotes+xml'):
            root=etree.fromstring(part.blob)
            for note in root.findall(qn('w:footnote')):
                if int(note.get(qn('w:id'),'0'))<=0:continue
                for p in note.iter(qn('w:p')):
                    pp=p.find(qn('w:pPr'))
                    if pp is None:pp=OxmlElement('w:pPr');p.insert(0,pp)
                    for tag in ['spacing','ind','jc']:
                        for el in list(pp.findall(qn('w:'+tag))):pp.remove(el)
                    spacing=OxmlElement('w:spacing')
                    for k,v in [('line','240'),('lineRule','auto'),('before','0'),('after','0')]:spacing.set(qn('w:'+k),v)
                    pp.append(spacing)
                    ind=OxmlElement('w:ind');ind.set(qn('w:firstLine'),'0');pp.append(ind)
                    jc=OxmlElement('w:jc');jc.set(qn('w:val'),'left');pp.append(jc)
                for run in note.iter(qn('w:r')):
                    rp=run.find(qn('w:rPr'))
                    if rp is None:rp=OxmlElement('w:rPr');run.insert(0,rp)
                    for tag in ['sz','szCs']:
                        for el in list(rp.findall(qn('w:'+tag))):rp.remove(el)
                        el=OxmlElement('w:'+tag);el.set(qn('w:val'),'20');rp.append(el)
                    if run.find(qn('w:footnoteRef')) is not None:
                        v=OxmlElement('w:vertAlign');v.set(qn('w:val'),'superscript');rp.append(v)
            part._blob=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
    for ref in doc.element.xpath('//w:footnoteReference'):
        rp=ref.getparent().get_or_add_rPr()
        v=OxmlElement('w:vertAlign');v.set(qn('w:val'),'superscript');rp.append(v)
    doc.core_properties.author=''; doc.core_properties.last_modified_by=''; doc.core_properties.comments=''
    # Enable review only after fields have been updated by Word (render script).
    revisions=doc.settings.element.find(qn('w:trackRevisions'))
    if revisions is not None: doc.settings.element.remove(revisions)
    doc.save(path)
    if len(numbered)!=len(manifest['labels']): raise RuntimeError('Not all labels became Word fields')

def main():
    p=argparse.ArgumentParser(); p.add_argument('--assemble-only',action='store_true'); args=p.parse_args()
    manifest=assemble()
    if not args.assemble_only:
        spec=importlib.util.spec_from_file_location('old_style',OLD/'tools/build_word.py'); module=importlib.util.module_from_spec(spec); spec.loader.exec_module(module)
        ref=ROOT/'tools/reference.docx'; module.reference_docx(ref)
        OUT.parent.mkdir(exist_ok=True)
        subprocess.run([str(PANDOC),str(ROOT/'report_for_word.md'),'--from=markdown+tex_math_dollars','--to=docx','--reference-doc='+str(ref),'--resource-path='+str(ROOT),'--output='+str(OUT)],check=True,cwd=ROOT)
        finish(OUT,manifest); manifest['output']=str(OUT); manifest['output_sha256']=sha(OUT)
        (ROOT/'build_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'main_words':manifest['main_words'],'counts':manifest['counts'],'literature':len(manifest['reference_keys']),'output':str(OUT)},ensure_ascii=False))

if __name__=='__main__': main()
